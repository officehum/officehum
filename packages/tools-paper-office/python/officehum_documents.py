"""The task-scoped document workspace API for submitted Python.

A run may read only the immutable version IDs its lease names, write only into the staging
directory, and declare the outputs it produced. Nothing here takes a filesystem path, a URL, a
file handle or a credential: a lease is not a directory, and this module is the whole of what it
exposes.

Typical use::

    import docx

    brief = documents.open_input("brief.docx")
    document = docx.Document(brief)
    document.add_paragraph("Reviewed by the office")
    document.save(documents.declare_output("reviewed.docx", "docx", edit_of="brief.docx"))
    documents.result({"paragraphs": len(document.paragraphs)})

Every refusal raises :class:`WorkspaceError`, whose ``code`` is a stable, inspectable reason that
the worker reports unchanged. The layout this module gives you is:

``<input>/<name>``    a read-only copy of a leased version, hashed again after the run
``<staging>/<name>``  the only writable directory in the workspace

Submitted code runs in a subprocess with a scrubbed environment, so ``documents.environment()``
never holds a credential. It is also honest about what it does not enforce: see the
``enforcement`` field the worker reports and ``README.md``.
"""

from __future__ import annotations

import json
import os
from typing import Any, Dict, Iterable, List, Optional

__all__ = [
    "WorkspaceError",
    "declare_output",
    "environment",
    "find_input",
    "inputs",
    "limits",
    "open_input",
    "outputs",
    "read_input",
    "result",
    "staging",
    "staging_directory",
    "write_output",
]

MAX_RESULT_BYTES = 64 * 1024
MAX_RESULT_DEPTH = 32
MAX_OUTPUTS = 256
OUTPUT_NAME_BYTES_MAX = 255


class WorkspaceError(RuntimeError):
    """A refusal. ``code`` is a stable reason the worker reports unchanged."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


class _Output:
    __slots__ = ("name", "format", "edit_of", "path")

    def __init__(self, name: str, format_: str, edit_of: Optional[str], path: str) -> None:
        self.name = name
        self.format = format_
        self.edit_of = edit_of
        self.path = path


class _Session:
    """One run's lease: the leased inputs, the staging directory, the permitted output formats."""

    def __init__(
        self,
        inputs: Dict[str, Dict[str, Any]],
        staging: str,
        output_formats: Iterable[str],
        limits: Dict[str, int],
        arguments: Optional[Dict[str, Any]] = None,
    ) -> None:
        self.inputs = inputs
        self.staging = staging
        self.output_formats = list(output_formats)
        self.limits = dict(limits)
        self.arguments = dict(arguments or {})
        self.outputs: Dict[str, _Output] = {}
        self.reads: List[str] = []
        self.result_value: Any = None
        self.result_declared = False

    def by_reference(self, reference: Any, what: str) -> Dict[str, Any]:
        if not isinstance(reference, str) or reference == "":
            raise WorkspaceError("ungranted-input", f"{what} must name a leased input version")
        if reference in self.inputs:
            return self.inputs[reference]
        matches = [entry for entry in self.inputs.values() if entry["name"] == reference]
        if len(matches) == 1:
            return matches[0]
        raise WorkspaceError(
            "ungranted-input",
            f"{what} names {reference!r}, which this task's lease does not grant; "
            "the lease lists: "
            + ", ".join(sorted(entry["name"] for entry in self.inputs.values())),
        )

    def declare(
        self, name: Any, format_: Any, edit_of: Optional[str], data: Optional[bytes]
    ) -> str:
        if not isinstance(name, str):
            raise WorkspaceError("traversal", "an output name must be a string")
        check_segment(name)
        if not isinstance(format_, str) or format_ not in self.output_formats:
            raise WorkspaceError(
                "output-format",
                f"this task may write only {', '.join(self.output_formats)}",
            )
        if not name.lower().endswith("." + format_):
            raise WorkspaceError(
                "output-format", f"an output named {name!r} is not a .{format_} file"
            )
        if name in self.outputs:
            raise WorkspaceError("invalid-request", f"{name!r} was already declared")
        if len(self.outputs) >= min(self.limits["fileCount"], MAX_OUTPUTS):
            raise WorkspaceError(
                "file-count",
                f"this task may write at most {min(self.limits['fileCount'], MAX_OUTPUTS)} files",
            )
        source: Optional[str] = None
        if edit_of is not None:
            entry = self.by_reference(edit_of, "edit_of")
            if not entry["writable"]:
                raise WorkspaceError(
                    "unleased-path",
                    f"this task's lease grants read-only access to {entry['name']!r}; "
                    "it cannot propose a new version of it",
                )
            source = entry["versionId"]
        path = os.path.join(self.staging, name)
        if os.path.lexists(path):
            raise WorkspaceError("invalid-request", f"{name!r} already exists in staging")
        self.outputs[name] = _Output(name, format_, source, path)
        if data is not None:
            try:
                with open(path, "xb") as handle:
                    handle.write(data)
            except OSError as error:
                if error.errno == 27:  # EFBIG: RLIMIT_FSIZE, the per-file allowance
                    raise WorkspaceError(
                        "file-size-limit", f"{name!r} exceeded the per-file size allowance"
                    ) from error
                raise WorkspaceError(
                    "invalid-request", f"{name!r} could not be written: {error}"
                ) from error
            except Exception as error:  # noqa: BLE001 - reported as a typed refusal
                raise WorkspaceError(
                    "invalid-request", f"{name!r} could not be written: {error}"
                ) from error
        return path

    def check_output_budget(self) -> None:
        total = 0
        for declared in self.outputs.values():
            try:
                total += os.lstat(declared.path).st_size
            except OSError as error:
                raise WorkspaceError(
                    "missing-output",
                    f"{declared.name!r} was declared but never written: {error}",
                ) from error
            if total > self.limits["outputBytes"]:
                raise WorkspaceError(
                    "output-bytes",
                    f"declared outputs exceed this task's {self.limits['outputBytes']} byte allowance",
                )


_SESSION: Optional[_Session] = None


def _configure(
    inputs: Dict[str, Dict[str, Any]],
    staging: str,
    output_formats: Iterable[str],
    limits: Dict[str, int],
    arguments: Optional[Dict[str, Any]] = None,
) -> _Session:
    """Bind one run's lease. Called by the worker before any submitted code runs."""
    global _SESSION
    _SESSION = _Session(inputs, staging, output_formats, limits, arguments)
    return _SESSION


def arguments() -> Dict[str, Any]:
    """The JSON arguments the host bound to this run, for code that is parameterised."""
    return dict(_session().arguments)


def _session() -> _Session:
    if _SESSION is None:
        raise WorkspaceError(
            "invalid-request", "the document workspace was not configured for this run"
        )
    return _SESSION


def check_segment(name: str) -> None:
    """A single path segment: no separators, no traversal, no control characters."""
    if (
        name in ("", ".", "..")
        or "/" in name
        or "\\" in name
        or len(name.encode("utf-8")) > OUTPUT_NAME_BYTES_MAX
        or any(ord(character) < 32 or ord(character) == 127 for character in name)
    ):
        raise WorkspaceError(
            "traversal",
            f"{name!r} is not a single file name; a lease grants no directory",
        )


def inputs() -> List[Dict[str, Any]]:
    """The leased inputs: version id, name, format, size, hash and whether they are writable."""
    session = _session()
    return [
        {
            "versionId": entry["versionId"],
            "fileId": entry["fileId"],
            "name": entry["name"],
            "format": entry["format"],
            "size": entry["size"],
            "sha256": entry["sha256"],
            "writable": entry["writable"],
        }
        for entry in sorted(session.inputs.values(), key=lambda item: item["name"])
    ]


def find_input(reference: str) -> Optional[str]:
    """The version id for a leased name or version id, or ``None`` when the lease has neither."""
    try:
        return _session().by_reference(reference, "reference")["versionId"]
    except WorkspaceError:
        return None


def open_input(reference: str) -> str:
    """The read-only path of a leased input, for a Paper library to open. Reads are recorded."""
    session = _session()
    entry = session.by_reference(reference, "open_input")
    if entry["versionId"] not in session.reads:
        session.reads.append(entry["versionId"])
    return entry["path"]


def read_input(reference: str) -> bytes:
    """A leased input's bytes. The input is read-only: writing to it would be refused afterwards."""
    session = _session()
    entry = session.by_reference(reference, "read_input")
    if entry["versionId"] not in session.reads:
        session.reads.append(entry["versionId"])
    with open(entry["path"], "rb") as handle:
        return handle.read()


def declare_output(name: str, format: str, edit_of: Optional[str] = None) -> str:
    """Reserve an output path in staging and return it, for a Paper library's ``save``."""
    return _session().declare(name, format, edit_of, None)


def write_output(
    name: str, data: bytes, format: Optional[str] = None, edit_of: Optional[str] = None
) -> str:
    """Write bytes to a declared output in one step, for data rather than a library's ``save``."""
    if format is None:
        format = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if not isinstance(data, (bytes, bytearray, memoryview)):
        raise WorkspaceError("invalid-request", "write_output takes bytes")
    return _session().declare(name, format, edit_of, bytes(data))


def outputs() -> List[Dict[str, Any]]:
    """The outputs declared so far, in declaration order."""
    return [
        {"name": declared.name, "format": declared.format, "editOf": declared.edit_of}
        for declared in _session().outputs.values()
    ]


def staging() -> str:
    """The one writable directory in the workspace."""
    return _session().staging


def staging_directory() -> str:
    """The one writable directory in the workspace (alias of :func:`staging`)."""
    return _session().staging


def limits() -> Dict[str, int]:
    """This task's resource budget, as the lease granted it."""
    return dict(_session().limits)


def environment() -> Dict[str, str]:
    """The worker's environment. It holds no credential: the host scrubs it before the spawn."""
    return dict(os.environ)


def result(value: Any) -> None:
    """Declare the run's JSON result. It is bounded, hashed into the receipt and shown to the model."""
    session = _session()
    try:
        encoded = json.dumps(value, allow_nan=False, sort_keys=True, default=None)
    except (TypeError, ValueError) as error:
        raise WorkspaceError("invalid-request", f"result is not JSON: {error}") from error
    if len(encoded.encode("utf-8")) > MAX_RESULT_BYTES:
        raise WorkspaceError(
            "invalid-request", f"result is larger than {MAX_RESULT_BYTES} bytes"
        )
    if _depth(value) > MAX_RESULT_DEPTH:
        raise WorkspaceError("invalid-request", "result is nested too deeply")
    session.result_value = json.loads(encoded)
    session.result_declared = True


def _depth(value: Any) -> int:
    if isinstance(value, dict):
        return 1 + max((_depth(item) for item in value.values()), default=0)
    if isinstance(value, (list, tuple)):
        return 1 + max((_depth(item) for item in value), default=0)
    return 1
