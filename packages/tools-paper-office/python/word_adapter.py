"""The typed Word adapter for the Paper Office worker (paper-docx).

One operation per call, six of them: ``inspect``, ``read``, ``create``, ``edit``, ``preview`` and
``compare``. Each takes a **leased input reference** — a FileStore version id with the name, hash and
size the lease pinned — never a server path, and answers with bounded JSON. Nothing here calls a model,
so the same operation and parameters produce the same result, and the worker that runs it binds the
code hash, the argument hash and the leased source hashes into one receipt.

The one thing worth stating plainly, because it looks like a contradiction: this module *does* handle
paths, but it never reads one from its parameters. ``documents.open_input(version_id)`` returns the
read-only copy of a leased version inside this run's input directory, and
``documents.declare_output(name, "docx", edit_of=...)`` returns a reserved path inside the run's
staging directory. The tool above this module takes file refs; the workspace API is what turns one into
a path, and it will only do that for a version the lease names.

What is preserved, and the evidence for it: an ``edit`` writes through ``docx.package.patch_save``,
which gives back the original bytes of every XML part semantically identical to its counterpart in the
source, so relationships, styles, headers, footers, numbering, theme and untouched sections survive the
open/edit/save cycle. The structural change set the result carries is
``docx.package.diff_package``'s own part-by-part comparison (with ``byteIdenticalCount`` as the
preservation witness) and the text diff is ``docx.package.text_diff``'s.

Bounded by construction: every page has a size, a cursor and a total; every string is truncated with an
explicit flag; the diff and the redline carry a line budget. A huge document cannot blow up a reply.

Refusals are ``documents.WorkspaceError`` with a stable code — ``document-protected``,
``malformed-package``, ``unsupported-structure``, ``boundary-violation``, ``target-not-found``,
``ambiguous-target`` — or the workspace's own codes for a bad name, a bad lease reference or a
read-only source. Every one of them leaves the source bytes unchanged: nothing here writes outside the
staging directory, and the workspace hashes the inputs again after the run.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import re
import zipfile
from typing import Any, Dict, List, NoReturn, Optional, Sequence, Tuple, Union

import docx
import docx.package as paper_package
from docx import blocks as blockops
from docx import commentops
from docx import protection
from docx.errors import (
    AmbiguousTargetError,
    BoundaryViolationError,
    DocumentProtectedError,
    MalformedPackageError,
    PaperRefusal,
    TargetNotFoundError,
)
from docx.search import find_one
from docx.story import iter_blocks, outline, story_parts
from docx.tableops import delete_row, find_table, insert_row_after, update_cell
from lxml.etree import XMLSyntaxError

import officehum_documents as documents

SCHEMA = "paper_word_operation"
VERSION = 1

MAX_PAGE = 50
DEFAULT_PAGE = 20
MAX_ITEM_TEXT = 400
MAX_LABEL_TEXT = 200
MAX_DIFF_LINES = 200
MAX_REDLINE_ENTRIES = 200
MAX_HEADINGS = 50
MAX_TABLES = 50
MAX_THREAD_ITEMS = 50
MAX_STEPS = 32
MAX_BLOCKS = 512
MAX_CELLS = 64
MAX_TEXT = 4_000

READ_SECTIONS = ("blocks", "paragraphs", "headings", "tables", "comments", "revisions")

BLOCK_KINDS = ("heading", "paragraph", "list", "table", "pageBreak")

STEP_FIELDS: Dict[str, Tuple[Tuple[str, ...], Tuple[str, ...]]] = {
    "insertSection": (
        ("anchor", "heading", "paragraphs"),
        ("headingStyle", "bodyStyle", "tracked", "author", "date"),
    ),
    "replaceParagraphs": (("anchor", "paragraphs"), ("count", "bodyStyle", "tracked", "author", "date")),
    "deleteParagraphs": (("anchor",), ("count", "tracked", "author", "date")),
    "replaceText": (("needle", "text"), ("story", "nth", "tracked", "author", "date")),
    "updateCell": (("nearText", "row", "column", "text"), ("tracked", "author", "date")),
    "insertRow": (("nearText", "row", "values"), ("copyFormatFrom",)),
    "deleteRow": (("nearText", "row"), ()),
    "addComment": (("anchor", "text"), ("author", "date")),
    "replyComment": (("commentId", "text"), ("author", "date")),
    "resolveComment": (("commentId",), ("resolved",)),
    "deleteComment": (("commentId",), ()),
    "acceptRevisions": ((), ("revisionAuthor",)),
    "rejectRevisions": ((), ("revisionAuthor",)),
}

REF_KEYS = ("fileId", "versionId", "name", "sha256", "size")
ANCHOR_TEXT_KEYS = ("text",)
ANCHOR_BLOCK_KEYS = ("story", "index", "contentHash")

HEADING_STYLE = re.compile(r"^Heading([1-9])$")
LIST_STYLE = {"bullet": "List Bullet", "decimal": "List Number"}


def _fail(code: str, message: str) -> NoReturn:
    """A typed refusal. The worker reports the code unchanged, and nothing was written."""
    raise documents.WorkspaceError(code, message)


# -- validation ------------------------------------------------------------------------------------


def _object(value: Any, what: str, allowed: Optional[Sequence[str]] = None) -> Dict[str, Any]:
    if not isinstance(value, dict):
        _fail("invalid-request", f"{what} must be a JSON object")
    if allowed is not None:
        for key in value:
            if key not in allowed:
                _fail("invalid-request", f"{what} carries an unknown field {key!r}")
    return value


def _required(entry: Dict[str, Any], keys: Sequence[str], what: str) -> None:
    for key in keys:
        if key not in entry:
            _fail("invalid-request", f"{what} is missing {key}")


def _string(value: Any, what: str, limit: int = MAX_TEXT) -> str:
    if not isinstance(value, str) or value == "":
        _fail("invalid-request", f"{what} must be a non-empty string")
    if len(value) > limit:
        _fail("invalid-request", f"{what} is longer than {limit} characters")
    return value


def _maybe_string(value: Any, what: str, limit: int = MAX_TEXT) -> str:
    if not isinstance(value, str):
        _fail("invalid-request", f"{what} must be a string, empty to delete the text")
    if len(value) > limit:
        _fail("invalid-request", f"{what} is longer than {limit} characters")
    return value


def _string_list(value: Any, what: str, limit: int = MAX_BLOCKS) -> List[str]:
    if not isinstance(value, list) or not value:
        _fail("invalid-request", f"{what} needs a non-empty list of strings")
    if len(value) > limit:
        _fail("invalid-request", f"{what} holds more than {limit} entries")
    return [_string(item, f"{what} entry") for item in value]


def _boolean(value: Any, what: str) -> bool:
    if not isinstance(value, bool):
        _fail("invalid-request", f"{what} must be true or false")
    return value


def _integer(value: Any, what: str, low: int, high: int) -> int:
    if not isinstance(value, int) or isinstance(value, bool) or not low <= value <= high:
        _fail("invalid-request", f"{what} must be an integer between {low} and {high}")
    return value


def _reference(value: Any, what: str) -> Dict[str, Any]:
    """A leased FileRef, as the lease described it. The workspace resolves the version id it names."""
    entry = _object(value, what, REF_KEYS)
    _required(entry, REF_KEYS, what)
    _string(entry["fileId"], f"{what} fileId", 256)
    _string(entry["versionId"], f"{what} versionId", 256)
    _string(entry["name"], f"{what} name", 255)
    _string(entry["sha256"], f"{what} sha256", 64)
    _integer(entry["size"], f"{what} size", 0, 2**53 - 1)
    return {key: entry[key] for key in REF_KEYS}


def _source(reference: Dict[str, Any]) -> Dict[str, Any]:
    return {key: reference[key] for key in REF_KEYS}


def _output_name(value: Any) -> str:
    name = _string(value, "an output name", 255)
    if not name.lower().endswith(".docx"):
        _fail("output-format", f"an output named {name!r} is not a .docx file")
    return name


def _author(value: Any, what: str = "author") -> str:
    return _string(value, what, 256)


def _date_of(value: Any, what: str) -> Optional[dt.datetime]:
    """An optional ISO 8601 instant, so a host can pin what a tracked change is stamped with."""
    if value is None:
        return None
    if not isinstance(value, str):
        _fail("invalid-request", f"{what} must be an ISO 8601 string")
    text = value[:-1] + "+00:00" if value.endswith("Z") else value
    try:
        parsed = dt.datetime.fromisoformat(text)
    except ValueError:
        _fail("invalid-request", f"{what} is not an ISO 8601 instant")
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=dt.timezone.utc)
    return parsed.astimezone(dt.timezone.utc)


def _truncated(value: Any, limit: int) -> Tuple[str, bool]:
    text = "" if value is None else str(value)
    if len(text) <= limit:
        return text, False
    return text[:limit], True


# -- opening a leased version ----------------------------------------------------------------------


def _leased_path(reference: Dict[str, Any]) -> str:
    """The read-only copy of a leased version. Reading it is recorded in the receipt's lineage."""
    return documents.open_input(reference["versionId"])


def _readable(reference: Dict[str, Any], path: str) -> None:
    """Structure triage before any XML is parsed, so a hostile package never reaches the parser."""
    diagnosis = paper_package.diagnose(path)
    if not diagnosis.readable:
        problems = "; ".join(diagnosis.problems) or "no detail"
        _fail(
            "malformed-package",
            f"the leased input {reference['name']!r} is not a readable Word package "
            f"({diagnosis.kind}): {problems}",
        )


def _load(what: str, path: str) -> "docx.document.Document":
    try:
        return docx.Document(path)
    except XMLSyntaxError as error:
        _fail("malformed-package", f"{what} holds XML that will not parse: {error}")
    except (PaperRefusal, ValueError, KeyError, zipfile.BadZipFile, UnicodeDecodeError) as error:
        _fail(
            "malformed-package",
            f"{what} is not a readable Word package: {type(error).__name__}: {error}",
        )


def _open(reference: Dict[str, Any]) -> Tuple[str, "docx.document.Document"]:
    """(path, document) for a leased version, or a typed refusal."""
    path = _leased_path(reference)
    what = f"the leased input {reference['name']!r}"
    _readable(reference, path)
    return path, _load(what, path)


def _reopen(path: str) -> "docx.document.Document":
    return _load("the written document", path)


def _applying(call, *args: Any, **kwargs: Any) -> Any:
    """Run one fork operation, turning its typed refusal into one of this workspace's codes."""
    try:
        return call(*args, **kwargs)
    except documents.WorkspaceError:
        raise
    except DocumentProtectedError as error:
        _fail("document-protected", str(error))
    except MalformedPackageError as error:
        _fail("malformed-package", str(error))
    except AmbiguousTargetError as error:
        _fail("ambiguous-target", str(error))
    except TargetNotFoundError as error:
        _fail("target-not-found", str(error))
    except BoundaryViolationError as error:
        _fail("boundary-violation", str(error))
    except PaperRefusal as error:
        _fail("unsupported-structure", f"{type(error).__name__}: {error}")


# -- traversal -------------------------------------------------------------------------------------


def _heading_level(style_id: Optional[str]) -> Optional[int]:
    if style_id is None:
        return None
    if style_id == "Title":
        return 0
    match = HEADING_STYLE.match(style_id)
    return int(match.group(1)) if match else None


class _View:
    """One traversal pass, cached: the blocks, their kinds, and the plain counts an inspection wants."""

    def __init__(self, document: "docx.document.Document") -> None:
        self.story_parts: Tuple[str, ...] = tuple(story_parts(document))
        self.blocks = tuple(iter_blocks(document, view="current"))
        self.paragraphs = tuple(block for block in self.blocks if block.kind == "paragraph")
        self.tables = tuple(block for block in self.blocks if block.kind == "table")
        self.headings = tuple(
            block for block in self.paragraphs if _heading_level(block.style_id) is not None
        )
        self.blind: Dict[str, int] = dict(sorted(outline(document).blind_region_counts.items()))

    def counts(self) -> Dict[str, int]:
        return {
            "storyParts": len(self.story_parts),
            "blocks": len(self.blocks),
            "paragraphs": len(self.paragraphs),
            "headings": len(self.headings),
            "tables": len(self.tables),
            "characters": sum(len(block.text) for block in self.blocks),
        }


def _anchor_of(block: Any) -> Dict[str, Any]:
    return {"story": block.story, "index": block.index, "contentHash": block.anchor.content_hash}


def _block_item(block: Any) -> Dict[str, Any]:
    text, truncated = _truncated(block.text, MAX_ITEM_TEXT)
    return {
        "story": block.story,
        "kind": block.kind,
        "index": block.index,
        "anchor": _anchor_of(block),
        "text": text,
        "truncated": truncated,
        "styleId": block.style_id,
        "headingLevel": _heading_level(block.style_id),
        "inInsert": block.in_insert,
        "inDelete": block.in_delete,
        "inContentControl": block.in_content_control,
        "inTextBox": block.in_text_box,
        "hasField": block.has_field,
        "table": block.table.to_dict() if block.table else None,
    }


def _cells_of(document: "docx.document.Document", block: Any) -> Optional[List[List[str]]]:
    """Cell text for one table block, matched to its live table by element identity."""
    for table in document.tables:
        if table._tbl is block._element:
            return [[_truncated(cell.text, MAX_ITEM_TEXT)[0] for cell in row.cells] for row in table.rows]
    return None


def _table_item(document: "docx.document.Document", block: Any, cells: bool) -> Dict[str, Any]:
    item = _block_item(block)
    shape = block.table
    item["shape"] = shape.to_dict() if shape else None
    item["cells"] = _cells_of(document, block) if cells else None
    return item


def _revision_item(revision: Any) -> Dict[str, Any]:
    text, truncated = _truncated(revision.text, MAX_ITEM_TEXT)
    return {
        "revisionType": revision.revision_type,
        "author": revision.author,
        "date": revision.date.isoformat() if revision.date else None,
        "text": text,
        "truncated": truncated,
        "story": revision.story,
        "contentHash": revision.anchor.content_hash,
        "isParagraphMark": revision.is_paragraph_mark,
        "resolvable": revision.is_resolvable,
    }


def _thread_items(thread: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    items: List[Dict[str, Any]] = []
    for entry in thread:
        text, truncated = _truncated(entry.get("text"), MAX_ITEM_TEXT)
        anchored, anchored_truncated = _truncated(entry.get("anchored_text"), MAX_LABEL_TEXT)
        items.append(
            {
                "commentId": entry.get("comment_id"),
                "author": entry.get("author"),
                "text": text,
                "truncated": truncated,
                "resolved": bool(entry.get("resolved")),
                "parentId": entry.get("parent_id"),
                "anchoredText": anchored,
                "anchoredTruncated": anchored_truncated,
            }
        )
    return items


def _page(items: Sequence[Any], cursor: int, page_size: int, section: str) -> Dict[str, Any]:
    total = len(items)
    if cursor > total:
        _fail(
            "invalid-request",
            f"the cursor {cursor} is past the end of the {section} page set, which holds {total} items",
        )
    window = list(items[cursor : cursor + page_size])
    returned = len(window)
    next_cursor = cursor + returned if cursor + returned < total else None
    return {
        "section": section,
        "pageSize": page_size,
        "cursor": cursor,
        "returned": returned,
        "total": total,
        "nextCursor": next_cursor,
        "hasMore": next_cursor is not None,
        "items": window,
    }


# -- diff, redline and change-set shaping -----------------------------------------------------------


def _text_diff(diff: Any) -> Dict[str, Any]:
    """Bounded, camelCased view of `paper-docx`'s per-story unified text diff."""
    raw = diff.to_dict()
    budget = MAX_DIFF_LINES
    stories: List[Dict[str, Any]] = []
    for story in raw["stories"]:
        lines = list(story["diff_lines"])
        kept = lines[: max(budget, 0)]
        budget -= len(kept)
        stories.append(
            {"story": story["story"], "diffLines": kept, "omittedLines": len(lines) - len(kept)}
        )
    return {
        "schema": raw["schema"],
        "version": raw["version"],
        "changedLineCount": raw["changed_line_count"],
        "stories": stories,
        "truncated": any(story["omittedLines"] > 0 for story in stories),
    }


def _change_set(diff: Any) -> Dict[str, Any]:
    """The structural change set: which parts were added, removed, changed or left byte-identical."""
    raw = diff.to_dict()
    return {
        "schema": raw["schema"],
        "version": raw["version"],
        "added": list(raw["added"]),
        "removed": list(raw["removed"]),
        "changed": [
            {
                "part": item["part"],
                "kind": item["kind"],
                "beforeSha256": item["before_sha256"],
                "afterSha256": item["after_sha256"],
                "semanticChange": item["semantic_change"],
            }
            for item in raw["changed"]
        ],
        "byteIdenticalCount": raw["byte_identical_count"],
        "semanticallyEmpty": raw["is_semantically_empty"],
    }


def _redline(path: str) -> Dict[str, Any]:
    """The tracked-change view of a saved document, plus what accepting every revision would do."""
    document = _reopen(path)
    entries: List[Dict[str, Any]] = []
    insertions = 0
    deletions = 0
    for block in iter_blocks(document, view="all"):
        if block.in_insert:
            insertions += 1
        if block.in_delete:
            deletions += 1
        if not (block.in_insert or block.in_delete):
            continue
        text, truncated = _truncated(block.text, MAX_ITEM_TEXT)
        marker = "insert+delete" if block.in_insert and block.in_delete else (
            "insert" if block.in_insert else "delete"
        )
        entries.append(
            {
                "story": block.story,
                "kind": block.kind,
                "index": block.index,
                "marker": marker,
                "text": text,
                "truncated": truncated,
            }
        )
    pending = _text_diff(_applying(paper_package.pending_changes, path))
    pending_lines: List[str] = []
    for story in pending["stories"]:
        pending_lines.extend(story["diffLines"])
    return {
        "schema": "paper_word_redline",
        "version": 1,
        "insertedBlocks": insertions,
        "deletedBlocks": deletions,
        "entries": entries[:MAX_REDLINE_ENTRIES],
        "total": len(entries),
        "truncated": len(entries) > MAX_REDLINE_ENTRIES,
        "pendingChangedLineCount": pending["changedLineCount"],
        "pendingLines": pending_lines[:MAX_DIFF_LINES],
    }


def _written(path: str) -> Tuple[str, int]:
    with open(path, "rb") as handle:
        data = handle.read()
    return hashlib.sha256(data).hexdigest(), len(data)


def _reopened(path: str) -> Dict[str, Any]:
    """Reopen the written document with the pinned library: the check the design wants before a commit."""
    document = _reopen(path)
    view = _View(document)
    return {
        "paragraphs": len(view.paragraphs),
        "headings": len(view.headings),
        "tables": len(view.tables),
        "comments": len(document.comments),
        "revisions": len(document.revisions),
        "storyParts": list(view.story_parts),
    }


def _envelope(operation: str, **fields: Any) -> Dict[str, Any]:
    return {"schema": SCHEMA, "version": VERSION, "operation": operation, "ok": True, **fields}


# -- operations ------------------------------------------------------------------------------------


def _inspect(params: Dict[str, Any]) -> Dict[str, Any]:
    entry = _object(params, "inspect parameters", ("ref",))
    _required(entry, ("ref",), "inspect parameters")
    reference = _reference(entry["ref"], "inspect ref")
    _path, document = _open(reference)
    view = _View(document)

    census: Dict[str, int] = {}
    resolvable = 0
    for revision in document.revisions:
        census[revision.revision_type] = census.get(revision.revision_type, 0) + 1
        if revision.is_resolvable:
            resolvable += 1
    thread = _applying(commentops.comment_thread, document)
    status = protection.protection_status(document)
    headings = [_block_item(block) for block in view.headings[:MAX_HEADINGS]]
    tables = [_table_item(document, block, False) for block in view.tables[:MAX_TABLES]]
    return _envelope(
        "inspect",
        source=_source(reference),
        document={
            **view.counts(),
            "comments": len(thread),
            "unresolvedComments": sum(1 for item in thread if not item["resolved"]),
            "threads": len({item.get("parent_id") or item.get("comment_id") for item in thread}),
            "revisions": len(document.revisions),
            "revisionTypes": dict(sorted(census.items())),
            "resolvableRevisions": resolvable,
            "unsupportedRevisions": _applying(document.revisions.remaining_unsupported),
            "blindRegions": view.blind,
            "protection": status.to_dict(),
            "editsBlocked": status.blocks_paper_edits,
        },
        headingOutline=headings,
        headingOutlineTruncated=len(view.headings) > MAX_HEADINGS,
        tables=tables,
        tablesTruncated=len(view.tables) > MAX_TABLES,
        comments=_thread_items(thread)[:MAX_THREAD_ITEMS],
        commentsTruncated=len(thread) > MAX_THREAD_ITEMS,
    )


def _section_items(
    document: "docx.document.Document", view: _View, section: str, cells: bool
) -> List[Dict[str, Any]]:
    if section == "blocks":
        return [_block_item(block) for block in view.blocks]
    if section == "paragraphs":
        return [_block_item(block) for block in view.paragraphs]
    if section == "headings":
        return [_block_item(block) for block in view.headings]
    if section == "tables":
        return [_table_item(document, block, cells) for block in view.tables]
    if section == "comments":
        return _thread_items(_applying(commentops.comment_thread, document))
    return [_revision_item(revision) for revision in document.revisions]


def _read(params: Dict[str, Any]) -> Dict[str, Any]:
    entry = _object(params, "read parameters", ("ref", "section", "cursor", "pageSize", "includeCells"))
    _required(entry, ("ref", "section"), "read parameters")
    reference = _reference(entry["ref"], "read ref")
    section = entry["section"]
    if section not in READ_SECTIONS:
        _fail("invalid-request", f"a read section must be one of {', '.join(READ_SECTIONS)}")
    cursor = 0 if entry.get("cursor") is None else _integer(entry["cursor"], "cursor", 0, 10**9)
    page_size = (
        DEFAULT_PAGE
        if entry.get("pageSize") is None
        else _integer(entry["pageSize"], "pageSize", 1, MAX_PAGE)
    )
    _path, document = _open(reference)
    view = _View(document)
    page = _page(
        _section_items(document, view, section, entry.get("includeCells") is True),
        cursor,
        page_size,
        section,
    )
    return _envelope("read", source=_source(reference), document=view.counts(), page=page)


def _block_specs(value: Any) -> List[Dict[str, Any]]:
    if not isinstance(value, list) or not value:
        _fail("invalid-request", "create needs a non-empty list of blocks")
    if len(value) > MAX_BLOCKS:
        _fail("invalid-request", f"create takes at most {MAX_BLOCKS} blocks")
    specs: List[Dict[str, Any]] = []
    for index, raw in enumerate(value):
        what = f"block {index}"
        spec = _object(raw, what)
        kind = _string(spec.get("kind"), f"{what} kind", 64)
        if kind not in BLOCK_KINDS:
            _fail("invalid-request", f"{what} names an unknown kind {kind!r}")
        if kind == "heading":
            _object(spec, what, ("kind", "text", "level"))
            _required(spec, ("kind", "text", "level"), what)
            _string(spec["text"], f"{what} text")
            _integer(spec["level"], f"{what} level", 0, 9)
        elif kind == "paragraph":
            _object(spec, what, ("kind", "text", "style"))
            _required(spec, ("kind", "text"), what)
            _string(spec["text"], f"{what} text")
            if spec.get("style") is not None:
                _string(spec["style"], f"{what} style", 255)
        elif kind == "list":
            _object(spec, what, ("kind", "items", "listKind", "level"))
            _required(spec, ("kind", "items"), what)
            _string_list(spec["items"], f"{what} items", MAX_CELLS)
            if spec.get("listKind", "bullet") not in ("bullet", "decimal"):
                _fail("invalid-request", f"{what} listKind must be bullet or decimal")
            if spec.get("level") is not None:
                _integer(spec["level"], f"{what} level", 0, 2)
        elif kind == "table":
            _object(spec, what, ("kind", "rows", "style"))
            _required(spec, ("kind", "rows"), what)
            rows = spec["rows"]
            if not isinstance(rows, list) or not rows:
                _fail("invalid-request", f"{what} needs a non-empty list of rows")
            if len(rows) > MAX_CELLS:
                _fail("invalid-request", f"{what} holds more than {MAX_CELLS} rows")
            width = None
            for row in rows:
                if not isinstance(row, list) or not row:
                    _fail("invalid-request", f"{what} rows must be non-empty lists of cells")
                if width is None:
                    width = len(row)
                if len(row) != width:
                    _fail("invalid-request", f"{what} rows must all hold the same number of cells")
                if len(row) > MAX_CELLS:
                    _fail("invalid-request", f"{what} rows hold more than {MAX_CELLS} cells")
                for cell in row:
                    if not isinstance(cell, str):
                        _fail("invalid-request", f"{what} cells must be strings")
                    if len(cell) > MAX_TEXT:
                        _fail("invalid-request", f"{what} holds a cell longer than {MAX_TEXT}")
            if spec.get("style") is not None:
                _string(spec["style"], f"{what} style", 255)
        else:
            _object(spec, what, ("kind",))
        specs.append(dict(spec))
    return specs


def _style(document: "docx.document.Document", name: str) -> str:
    """The document's own name for a style, so `add_paragraph` is not handed a style id."""
    for style in document.styles:
        if style.name == name or style.style_id == name:
            return style.name
    _fail("unsupported-structure", f"this document defines no style named {name!r}")


def _add_block(document: "docx.document.Document", spec: Dict[str, Any]) -> None:
    kind = spec["kind"]
    if kind == "heading":
        level = spec["level"]
        document.add_paragraph(
            spec["text"], style=_style(document, "Title" if level == 0 else f"Heading {level}")
        )
    elif kind == "paragraph":
        document.add_paragraph(
            spec["text"],
            style=None if spec.get("style") is None else _style(document, spec["style"]),
        )
    elif kind == "list":
        level = spec.get("level", 0)
        style = _style(document, LIST_STYLE[spec.get("listKind", "bullet")] + ("" if level == 0 else f" {level + 1}"))
        for item in spec["items"]:
            document.add_paragraph(item, style=style)
    elif kind == "table":
        rows = spec["rows"]
        table = document.add_table(
            rows=len(rows),
            cols=len(rows[0]),
            style=None if spec.get("style") is None else _style(document, spec["style"]),
        )
        for row_index, row in enumerate(rows):
            for column_index, cell in enumerate(row):
                table.cell(row_index, column_index).text = cell
    else:
        document.add_page_break()


def _create(params: Dict[str, Any]) -> Dict[str, Any]:
    entry = _object(params, "create parameters", ("output", "blocks"))
    _required(entry, ("output", "blocks"), "create parameters")
    name = _output_name(entry["output"])
    specs = _block_specs(entry["blocks"])
    path = documents.declare_output(name, "docx")
    document = docx.Document()
    for spec in specs:
        _applying(_add_block, document, spec)
    _applying(document.save, path)
    sha, size = _written(path)
    return _envelope(
        "create",
        draft={
            "status": "draft",
            "name": name,
            "editOf": None,
            "sha256": sha,
            "size": size,
            "published": False,
        },
        structure=_View(document).counts(),
        provenance={"template": "paper-docx default", "source": None},
        reopened=_reopened(path),
    )


# -- edits -----------------------------------------------------------------------------------------


def _anchor(document: "docx.document.Document", value: Any, what: str) -> Union[str, Any]:
    """A live mutation anchor: an exact string the fork resolves with `find_one`, or a reacquired block.

    A block anchor is reacquired from the current projection and checked against the content hash the
    caller read earlier, so a stale index refuses instead of editing whatever has moved into its place.
    """
    spec = _object(value, what)
    if "text" in spec:
        _object(spec, what, ANCHOR_TEXT_KEYS)
        return _string(spec["text"], f"{what} text")
    _object(spec, what, ANCHOR_BLOCK_KEYS)
    _required(spec, ANCHOR_BLOCK_KEYS, what)
    story = _string(spec["story"], f"{what} story", 255)
    index = _integer(spec["index"], f"{what} index", 0, 10**6)
    content_hash = _string(spec["contentHash"], f"{what} contentHash", 64)
    matches = [
        block
        for block in iter_blocks(document, view="current")
        if block.story == story and block.index == index
    ]
    if len(matches) != 1:
        _fail(
            "target-not-found",
            f"{what} names {story} block {index}, which this document does not hold exactly once",
        )
    block = matches[0]
    if block.anchor.content_hash != content_hash:
        _fail(
            "target-not-found",
            f"{what} names {story} block {index}, whose content no longer matches {content_hash}; "
            "read the document again and edit what is there now",
        )
    return block


def _step_fields(spec: Dict[str, Any], kind: str, what: str) -> Dict[str, Any]:
    required, optional = STEP_FIELDS[kind]
    _object(spec, what, ("kind", *required, *optional))
    _required(spec, required, what)
    return spec


def _step_tracked(spec: Dict[str, Any], default: bool = True) -> bool:
    return _boolean(spec.get("tracked", default), "tracked")


def _count_of(spec: Dict[str, Any], what: str) -> int:
    return _integer(spec.get("count", 1), f"{what} count", 1, 1_000)


def _comment_of(document: "docx.document.Document", spec: Dict[str, Any], what: str) -> Any:
    comment_id = _integer(spec["commentId"], f"{what} commentId", 0, 10**6)
    comment = document.comments.get(comment_id)
    if comment is None:
        _fail("target-not-found", f"{what} names comment {comment_id}, which this document does not hold")
    return comment_id, comment


def _apply(
    document: "docx.document.Document",
    spec: Any,
    plan_author: str,
    plan_date: Optional[dt.datetime],
) -> Dict[str, Any]:
    """One typed edit step, against the fork's own primitives. Returns the fork's change evidence."""
    block = _object(spec, "an edit step")
    kind = _string(block.get("kind"), "an edit step kind", 64)
    if kind not in STEP_FIELDS:
        _fail("invalid-request", f"an edit step names an unknown kind {kind!r}")
    what = f"the {kind} step"
    fields = _step_fields(block, kind, what)
    author = _author(fields.get("author", plan_author), f"{what} author")
    date = _date_of(fields.get("date"), f"{what} date") or plan_date

    if kind == "insertSection":
        anchor = _anchor(document, fields["anchor"], f"{what} anchor")
        tracked = _step_tracked(fields)
        result = _applying(
            blockops.insert_section_after,
            document,
            anchor,
            heading=_string(fields["heading"], f"{what} heading"),
            paragraphs=_string_list(fields["paragraphs"], f"{what} paragraphs"),
            heading_style=_string(fields.get("headingStyle", "Heading2"), f"{what} headingStyle", 255),
            body_style=fields.get("bodyStyle"),
            tracked=tracked,
            author=author,
            date=date,
        )
        return {"kind": kind, "tracked": tracked, "evidence": result.to_dict()}

    if kind in ("replaceParagraphs", "deleteParagraphs"):
        if not _step_tracked(fields):
            _fail(
                "unsupported-structure",
                f"{what} cannot run untracked: paper-docx removes and replaces paragraphs only as "
                "tracked changes, so this edit is always a Word revision",
            )
        anchor = _anchor(document, fields["anchor"], f"{what} anchor")
        count = _count_of(fields, what)
        if kind == "replaceParagraphs":
            result = _applying(
                blockops.tracked_replace_paragraphs,
                document,
                anchor,
                _string_list(fields["paragraphs"], f"{what} paragraphs"),
                count=count,
                body_style=fields.get("bodyStyle"),
                author=author,
                date=date,
            )
        else:
            result = _applying(
                blockops.tracked_delete_paragraphs, document, anchor, count=count, author=author, date=date
            )
        return {"kind": kind, "tracked": True, "evidence": result.to_dict()}

    if kind == "replaceText":
        span = _applying(
            find_one,
            document,
            _string(fields["needle"], f"{what} needle"),
            nth=None if fields.get("nth") is None else _integer(fields["nth"], f"{what} nth", 1, 10**6),
            story=fields.get("story"),
            view="current",
            match="exact",
        )
        tracked = _step_tracked(fields)
        result = _applying(
            span.replace,
            _maybe_string(fields["text"], f"{what} text"),
            tracked=tracked,
            author=author,
            date=date,
        )
        return {"kind": kind, "tracked": tracked, "evidence": result.to_dict()}

    if kind == "updateCell":
        table = _applying(
            find_table, document, near_text=_string(fields["nearText"], f"{what} nearText")
        )
        tracked = _step_tracked(fields)
        result = _applying(
            update_cell,
            table,
            _integer(fields["row"], f"{what} row", 0, 10**6),
            _integer(fields["column"], f"{what} column", 0, 10**6),
            _maybe_string(fields["text"], f"{what} text"),
            tracked=tracked,
            author=author,
            date=date,
        )
        return {"kind": kind, "tracked": tracked, "evidence": result.to_dict()}

    if kind == "insertRow":
        table = _applying(
            find_table, document, near_text=_string(fields["nearText"], f"{what} nearText")
        )
        values = _string_list(fields["values"], f"{what} values", MAX_CELLS)
        _applying(
            insert_row_after,
            table,
            _integer(fields["row"], f"{what} row", 0, 10**6),
            values,
            copy_format_from=None
            if fields.get("copyFormatFrom") is None
            else _integer(fields["copyFormatFrom"], f"{what} copyFormatFrom", 0, 10**6),
        )
        return {
            "kind": kind,
            "tracked": False,
            "evidence": {"insertedValues": values, "note": "paper-docx has no tracked row insertion"},
        }

    if kind == "deleteRow":
        table = _applying(
            find_table, document, near_text=_string(fields["nearText"], f"{what} nearText")
        )
        row = _integer(fields["row"], f"{what} row", 0, 10**6)
        _applying(delete_row, table, row)
        return {
            "kind": kind,
            "tracked": False,
            "evidence": {"deletedRow": row, "note": "paper-docx has no tracked row deletion"},
        }

    if kind == "addComment":
        anchor = _anchor(document, fields["anchor"], f"{what} anchor")
        if isinstance(anchor, str):
            anchor = _applying(find_one, document, anchor, view="current", match="exact")
        comment = _applying(
            anchor.comment, _string(fields["text"], f"{what} text"), author=author, date=date
        )
        return {
            "kind": kind,
            "tracked": False,
            "evidence": {"commentId": comment.comment_id, "author": author},
        }

    if kind in ("replyComment", "resolveComment", "deleteComment"):
        comment_id, comment = _comment_of(document, fields, what)
        if kind == "replyComment":
            reply = _applying(
                commentops.reply,
                document,
                comment,
                _string(fields["text"], f"{what} text"),
                author=author,
                date=date,
            )
            return {
                "kind": kind,
                "tracked": False,
                "evidence": {"commentId": reply.comment_id, "parentId": comment_id, "author": author},
            }
        if kind == "resolveComment":
            resolved = True if fields.get("resolved") is None else _boolean(fields["resolved"], f"{what} resolved")
            _applying(commentops.resolve, document, comment, resolved=resolved)
            return {
                "kind": kind,
                "tracked": False,
                "evidence": {"commentId": comment_id, "resolved": resolved},
            }
        _applying(commentops.delete_comment, document, comment)
        return {"kind": kind, "tracked": False, "evidence": {"deletedCommentId": comment_id}}

    revisions = document.revisions
    accept = kind == "acceptRevisions"
    resolved = _applying(
        revisions.accept_all if accept else revisions.reject_all,
        author=None if fields.get("revisionAuthor") is None else _author(fields["revisionAuthor"], f"{what} revisionAuthor"),
    )
    return {
        "kind": kind,
        "tracked": False,
        "evidence": {
            "resolved": resolved,
            "accepted": accept,
            "remainingUnsupported": _applying(document.revisions.remaining_unsupported),
        },
    }


def _steps(value: Any) -> List[Dict[str, Any]]:
    if not isinstance(value, list) or not value:
        _fail("invalid-request", "edit needs a non-empty list of steps")
    if len(value) > MAX_STEPS:
        _fail("invalid-request", f"edit takes at most {MAX_STEPS} steps")
    return [_object(step, f"step {index}") for index, step in enumerate(value)]


def _edit(params: Dict[str, Any]) -> Dict[str, Any]:
    entry = _object(
        params,
        "edit parameters",
        ("ref", "output", "steps", "author", "date", "acknowledgeProtection"),
    )
    _required(entry, ("ref", "output", "steps", "author"), "edit parameters")
    reference = _reference(entry["ref"], "edit source")
    name = _output_name(entry["output"])
    plan_author = _author(entry["author"], "edit author")
    plan_date = _date_of(entry.get("date"), "edit date")
    steps = _steps(entry["steps"])

    path = documents.declare_output(name, "docx", edit_of=reference["versionId"])
    source_path, document = _open(reference)
    status = protection.protection_status(document)
    if status.blocks_paper_edits:
        if entry.get("acknowledgeProtection") is not True:
            _fail(
                "document-protected",
                f"the source declares document protection (edit={status.edit!r}, enforced), which "
                "paper-docx refuses to edit. Review why it is locked, then pass acknowledgeProtection "
                "to proceed deliberately",
            )
        _applying(protection.acknowledge_protection, document)

    changes: List[Dict[str, Any]] = []
    for spec in steps:
        changes.append(_apply(document, spec, plan_author, plan_date))

    patch = _applying(paper_package.patch_save, source_path, document, path)
    change_set = _applying(paper_package.diff_package, source_path, path)
    text_diff = _applying(paper_package.text_diff, source_path, path)
    sha, size = _written(path)
    return _envelope(
        "edit",
        source=_source(reference),
        proposal={
            "status": "proposed",
            "name": name,
            "editOf": reference["versionId"],
            "sha256": sha,
            "size": size,
            "published": False,
        },
        changes=changes,
        changeSet=_change_set(change_set),
        diff=_text_diff(text_diff),
        redline=_redline(path),
        preservation={
            "restoredParts": list(patch.restored_parts),
            "changedParts": list(patch.changed_parts),
            "addedParts": list(patch.added_parts),
            "removedParts": list(patch.removed_parts),
            "verbatimCopy": patch.verbatim_copy,
        },
        protection=protection.protection_status(document).to_dict(),
        reopened=_reopened(path),
    )


def _label(block: Any) -> str:
    """One bounded preview line, with the structure marker where a structure exists."""
    if block.kind == "table":
        text, _truncated_text = _truncated(block.text.replace("\n", " | "), MAX_LABEL_TEXT)
        shape = block.table
        marker = f"T{shape.rows}x{shape.columns}" if shape else "T"
        return f"{marker}: {text}"
    level = _heading_level(block.style_id)
    text, _truncated_text = _truncated(block.text, MAX_LABEL_TEXT)
    if level is not None:
        return f"H{level}: {text}"
    if (block.style_id or "").startswith("List"):
        return f"L: {text}"
    return f"P: {text}"


def _preview(params: Dict[str, Any]) -> Dict[str, Any]:
    entry = _object(params, "preview parameters", ("ref", "cursor", "pageSize"))
    _required(entry, ("ref",), "preview parameters")
    reference = _reference(entry["ref"], "preview ref")
    cursor = 0 if entry.get("cursor") is None else _integer(entry["cursor"], "cursor", 0, 10**9)
    page_size = (
        DEFAULT_PAGE
        if entry.get("pageSize") is None
        else _integer(entry["pageSize"], "pageSize", 1, MAX_PAGE)
    )
    path, document = _open(reference)
    view = _View(document)
    items = [
        {
            "story": block.story,
            "kind": block.kind,
            "index": block.index,
            "label": _label(block),
            "truncated": len(block.text) > MAX_LABEL_TEXT,
        }
        for block in view.blocks
    ]
    thread = _applying(commentops.comment_thread, document)
    return _envelope(
        "preview",
        source=_source(reference),
        page=_page(items, cursor, page_size, "preview"),
        structure={
            **view.counts(),
            "comments": len(thread),
            "revisions": len(document.revisions),
        },
        protection=protection.protection_status(document).to_dict(),
        pending=_text_diff(_applying(paper_package.pending_changes, path)),
        redline=_redline(path),
    )


def _compare(params: Dict[str, Any]) -> Dict[str, Any]:
    entry = _object(params, "compare parameters", ("from", "to"))
    _required(entry, ("from", "to"), "compare parameters")
    from_ref = _reference(entry["from"], "compare from")
    to_ref = _reference(entry["to"], "compare to")
    from_path = _leased_path(from_ref)
    to_path = _leased_path(to_ref)
    _readable(from_ref, from_path)
    _readable(to_ref, to_path)
    change_set = _applying(paper_package.diff_package, from_path, to_path)
    text_diff = _applying(paper_package.text_diff, from_path, to_path)
    return _envelope(
        "compare",
        sources={"from": _source(from_ref), "to": _source(to_ref)},
        changeSet=_change_set(change_set),
        diff=_text_diff(text_diff),
        redline=_redline(to_path),
        identical=change_set.is_semantically_empty and text_diff.is_empty,
    )


OPERATIONS = {
    "inspect": _inspect,
    "read": _read,
    "create": _create,
    "edit": _edit,
    "preview": _preview,
    "compare": _compare,
}


def run(operation: str, parameters: Any) -> Dict[str, Any]:
    """Run one typed Word operation. ``operation`` is fixed by the submitted code, not by the caller."""
    if operation not in OPERATIONS:
        _fail("invalid-request", f"{operation!r} is not a Word operation")
    return OPERATIONS[operation](_object(parameters, "operation parameters"))


__all__ = ["run", "SCHEMA", "VERSION", "OPERATIONS"]